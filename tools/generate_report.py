from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

base=Path(__file__).resolve().parents[1]
out=base/'docs/B24DCCC047_PhamHaiDang_BaoCao.docx'
doc=Document(); section=doc.sections[0]
section.page_width=Inches(8.5); section.page_height=Inches(11)
section.top_margin=section.bottom_margin=Inches(.65)
section.left_margin=section.right_margin=Inches(.7)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Caption']:
 s=doc.styles[name]; s.font.name='DejaVu Sans'; s.font.color.rgb=RGBColor(0,0,0)
 s.font.size=Pt(11 if name=='Normal' else 10 if name=='Caption' else 19 if name=='Title' else 16 if name=='Heading 1' else 12)
 s.paragraph_format.space_after=Pt(8)
 s.paragraph_format.line_spacing=1.12
normal=doc.styles['Normal']; normal.paragraph_format.widow_control=True
doc.styles['Caption'].font.bold=False
for node in list(doc.styles.element.xpath('.//w:pBdr')):
 node.getparent().remove(node)

def p(t):
 para=doc.add_paragraph(t); para.paragraph_format.space_before=Pt(6)
 return para
def h(t): doc.add_heading(t,level=1)
def page(t): doc.add_page_break(); h(t)
def images(left,right,caption):
 para=doc.add_paragraph(); para.alignment=WD_ALIGN_PARAGRAPH.CENTER
 para.paragraph_format.space_after=Pt(3)
 para.add_run().add_picture(str(base/'docs/screenshots'/left),width=Inches(2.8))
 para.add_run('   ')
 para.add_run().add_picture(str(base/'docs/screenshots'/right),width=Inches(2.8))
 cp=doc.add_paragraph(caption,style='Caption'); cp.alignment=WD_ALIGN_PARAGRAPH.CENTER

def table(headers, rows):
 t=doc.add_table(rows=1,cols=len(headers)); t.style='Table Grid'
 t.autofit=False
 widths=[2,5.1] if headers[0]=='Thành phần' else [3.2,3.9]
 for col,width in zip(t.columns,widths): col.width=Inches(width)
 for c,text in zip(t.rows[0].cells,headers):
  c.text=text
  for run in c.paragraphs[0].runs: run.bold=True; run.font.color.rgb=RGBColor(0,0,0)
  pr=c._tc.get_or_add_tcPr(); shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E7E6E6');pr.append(shade)
 for row in rows:
  for c,text in zip(t.add_row().cells,row): c.text=text
 for row_index,row in enumerate(t.rows):
  for c,width in zip(row.cells,widths):
   c.width=Inches(width)
   c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   pr=c._tc.get_or_add_tcPr()
   borders=OxmlElement('w:tcBorders')
   for edge in ['top','left','bottom','right']:
    e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
   pr.append(borders)
   margins=OxmlElement('w:tcMar')
   for edge in ['top','bottom','left','right']:
    e=OxmlElement('w:'+edge);e.set(qn('w:w'),'80');e.set(qn('w:type'),'dxa');margins.append(e)
   pr.append(margins)
   if row_index>0:
    shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'FFFFFF');pr.append(shade)
   for para in c.paragraphs:
    para.paragraph_format.space_after=Pt(2)
    for run in para.runs: run.font.size=Pt(10)
 return t

# Snippets are read directly from the Java/XML sources, with verified line numbers.
# No duplicated source listings are maintained in the report.
import textwrap
java_dir = 'app/src/main/java/com/example/myapplication/'

def snippet(title, filename, start, end=None, include_end=False):
 source_path = java_dir + filename if filename.endswith('.java') else filename
 lines = (base / source_path).read_text().splitlines()
 first = next(i for i, line in enumerate(lines) if start in line)
 last = len(lines) - 1 if lines[-1].strip() == '}' else len(lines)
 if end is not None:
  last = next(i for i in range(first + 1, len(lines)) if end in lines[i])
  if include_end: last += 1
 while last > first and not lines[last-1].strip(): last -= 1
 if lines[first].startswith('        ') and lines[last-1] == '    }': last -= 1
 doc.add_heading(title, level=2)
 note = doc.add_paragraph(f'{filename} - dòng {first+1} đến {last}', style='Caption')
 note.paragraph_format.space_after=Pt(5)
 content = textwrap.dedent('\n'.join(lines[first:last]))
 for line in content.splitlines():
  para=doc.add_paragraph()
  para.paragraph_format.space_after=Pt(0)
  para.paragraph_format.space_before=Pt(0)
  para.paragraph_format.line_spacing=1.05
  para.paragraph_format.left_indent=Inches(.1)
  para.paragraph_format.keep_with_next=True
  run=para.add_run(line if line else ' ')
  run.font.name='DejaVu Sans Mono'; run.font.size=Pt(9)
 # Allow a break after, but never through, each source excerpt.
 para.paragraph_format.keep_with_next=False

# Overview
p0=doc.add_paragraph('BÁO CÁO LẬP TRÌNH ANDROID',style='Subtitle')
doc.add_paragraph('Ứng dụng quản lý danh bạ liên hệ',style='Title')
p('Sinh viên: Phạm Hải Đăng\nMã sinh viên: B24DCCC047\nLớp: D24CQPM03\nNgày hoàn thành: 28/10/2026')
p('Repository GitHub\nhttps://github.com/DANG-PH/B24DCCC047_PhamHaiDang_QuanLyDanhBa')
h('1 Mục tiêu và phạm vi')
p('Ứng dụng lưu và quản lý các liên hệ cá nhân bằng Java và XML. Bài làm gồm bốn Activity: đăng nhập, danh sách, thêm liên hệ và chi tiết. Chức năng chính gồm xác thực tài khoản cố định, kiểm tra dữ liệu nhập, đánh dấu yêu thích, xóa liên hệ và mở màn hình quay số.')
p('Dữ liệu nằm trong ArrayList<Contact> của ContactRepository. Khi tiến trình bị kết thúc, danh bạ được khởi tạo lại với ba liên hệ mẫu. Ứng dụng không dùng SQLite, Room, Firebase hoặc API; tài khoản thực hành là admin / 123456.')
h('2 Cấu trúc và luồng xử lý')
table(['Thành phần','Vai trò'],[
 ('Contact','Mô hình Serializable: id, họ tên, số điện thoại, nhóm, giới tính, yêu thích, ghi chú.'),
 ('ContactRepository','Giữ ArrayList, thêm không trùng số, tìm theo ID, xóa và đếm yêu thích.'),
 ('ContactValidator','Kiểm tra tên không trống và số điện thoại có định dạng 0[0-9]{9}.'),
 ('ContactAdapter','Hiển thị từng liên hệ trong ListView với avatar và dấu sao.'),
 ('BaseActivity','Áp dụng khoảng trống thanh hệ thống/bàn phím, dùng chung nút Quay lại.')])
p('Luồng màn hình: LoginActivity → MainActivity → AddContactActivity hoặc ContactDetailActivity. Đăng nhập thành công gọi finish() để Back không trở về đăng nhập. Các màn hình thêm và chi tiết gọi finish() để quay lại danh sách; MainActivity.onResume() nạp lại dữ liệu và thống kê.')

page('3 Màn hình đăng nhập LoginActivity')
p('Giao diện XML có EditText tên đăng nhập, EditText mật khẩu với inputType="textPassword", TextView báo lỗi và nút Đăng nhập. LoginActivity khai báo USERNAME="admin", PASSWORD="123456". Hàm login() chặn bỏ trống và hiển thị đúng thông báo khi sai tài khoản; nếu đúng, dùng Intent mở MainActivity rồi finish().')
images('01_login.png','02_login_error.png','Hình 1 bên trái: giao diện đăng nhập. Hình 2 bên phải: thông báo sai tài khoản.')
p('Kiểm thử đã thực hiện: bỏ trống cả hai trường, nhập admin với mật khẩu sai, sau đó nhập 123456. Hai trường hợp đầu ở lại đăng nhập; trường hợp đúng chuyển sang danh sách có 3 liên hệ và 1 yêu thích. Activity đăng nhập kết thúc sau khi chuyển màn hình.')

page('Mã tiêu biểu của màn hình đăng nhập')
p('Các trích đoạn sau được lấy trực tiếp từ mã nguồn. Tài khoản nằm trong LoginActivity; login() đọc hai trường nhập và quyết định hiển thị lỗi hay chuyển Activity.')
snippet('Tài khoản cố định', 'LoginActivity.java', 'private static final String USERNAME', '@Override protected void onCreate')
snippet('Kiểm tra dữ liệu và chuyển màn hình', 'LoginActivity.java', '        if (name.isEmpty()', None)
p('Nhánh đầu chặn thiếu thông tin; nhánh tiếp theo dùng đúng thông báo lỗi quy định. Chỉ nhánh else mở MainActivity và finish() màn hình đăng nhập. TextView báo lỗi ban đầu được đặt visibility="gone" trong XML.')

page('4 Danh sách liên hệ MainActivity')
p('MainActivity dùng ListView và ContactAdapter. Mỗi dòng gồm chữ cái đầu của tên gọi, họ tên, số điện thoại, nhóm và dấu sao khi yêu thích. Khối thống kê tính tổng số phần tử và số liên hệ yêu thích. Nút Thêm liên hệ mở AddContactActivity; nhấn một dòng truyền Contact Serializable qua Intent sang ContactDetailActivity.')
images('03_contacts.png','07_contacts_added.png','Hình 3 bên trái: dữ liệu mẫu. Hình 4 bên phải: danh sách sau khi thêm liên hệ yêu thích.')
p('Dữ liệu mẫu đúng đề: Nguyễn Văn An – 0901234567 – Bạn bè; Trần Thị Bình – 0912345678 – Gia đình – yêu thích; Lê Minh Châu – 0987654321 – Công việc. Sau khi thêm Phạm Hải Đăng, thống kê đổi từ 3 liên hệ/1 yêu thích thành 4 liên hệ/2 yêu thích. onResume() cập nhật lại adapter và hai TextView thống kê.')

page('Mã tiêu biểu của danh sách liên hệ')
snippet('Intent mở chi tiết và thêm liên hệ', 'MainActivity.java', '        list.setOnItemClickListener', '    @Override protected void onResume')
snippet('Nạp lại dữ liệu và thống kê khi quay lại', 'MainActivity.java', '    @Override protected void onResume', None)
snippet('Hiển thị avatar và dấu sao trong Adapter', 'ContactAdapter.java', '        Contact contact = getItem(position);', '        return view;')
p('Contact là Serializable nên putExtra() truyền được đối tượng qua Intent. Danh sách dùng bản sao ArrayList từ repository; dữ liệu gốc được dùng để cập nhật yêu thích. onResume() nạp lại dữ liệu sau thêm, xóa hoặc thay đổi yêu thích.')

page('5 Thêm liên hệ AddContactActivity')
p('Biểu mẫu XML gồm họ tên, điện thoại (inputType="phone"), Spinner với Gia đình/Bạn bè/Công việc/Khác, RadioGroup Nam/Nữ/Khác (mặc định Nam), CheckBox yêu thích và ghi chú textMultiLine. Nút Lưu gọi save(): kiểm tra tên, định dạng số, rồi kiểm tra trùng bằng containsPhone(). Lỗi được gắn vào EditText bằng setError().')
images('06_add_form.png','05_add_duplicate.png','Hình 5 bên trái: biểu mẫu điền hợp lệ. Hình 6 bên phải: từ chối số điện thoại đã có.')
p('Khi hợp lệ, save() tạo Contact từ toàn bộ trường nhập và gọi ContactRepository.add(). Sau đó hiển thị Toast, finish() về danh sách. Đã kiểm tra tên trống, số 123, số mẫu 0901234567 bị trùng và số mới 0966666666 lưu thành công; không có bản ghi mới khi dữ liệu sai. Ghi chú và trạng thái yêu thích được hiển thị lại trong chi tiết.')

page('Mã tiêu biểu của biểu mẫu thêm liên hệ')
snippet('Kiểm tra tên số điện thoại và số trùng', 'AddContactActivity.java', '        if (!ContactValidator.isValidName', '        Spinner group = findViewById(R.id.groupInput);')
snippet('Quy tắc số điện thoại', 'ContactValidator.java', '    public static boolean isValidPhone', None)
snippet('Tạo Contact và quay về danh sách', 'AddContactActivity.java', '        Contact contact = new Contact', None)
p('save() đọc tên, điện thoại, Spinner, RadioGroup, CheckBox và ghi chú trước khi tạo đối tượng. Nhánh lỗi gọi return nên không thêm dữ liệu. Hàm validator chỉ chấp nhận đúng 10 chữ số ASCII, chữ số đầu là 0; containsPhone() kiểm tra trên toàn bộ danh bạ.')

page('6 Chi tiết liên hệ ContactDetailActivity')
p('Activity nhận Contact bằng IntentCompat.getSerializableExtra(), tìm bản gốc trong repository theo ID rồi hiển thị avatar, tên, số, nhóm, giới tính và ghi chú. Nếu dữ liệu không còn, ứng dụng thông báo và đóng màn hình. CheckBox cập nhật trực tiếp favorite của liên hệ; nút Quay lại gọi finish().')
images('08_detail.png','09_dialer.png','Hình 7 bên trái: đầy đủ thông tin liên hệ mới. Hình 8 bên phải: trình quay số mở từ nút Gọi điện.')
p('Nút Gọi điện tạo Intent.ACTION_DIAL với Uri.fromParts("tel", số điện thoại, null). Ứng dụng mở trình quay số để người dùng quyết định gọi và không cần quyền CALL_PHONE. Kiểm thử xác nhận Intent có ACTION_DIAL và dữ liệu tel:0901234567; ảnh bên phải ghi lại trình quay số thực tế cho Nguyễn Văn An trong kịch bản này.')

page('Mã tiêu biểu của màn hình chi tiết')
snippet('Nhận liên hệ qua Intent', 'ContactDetailActivity.java', '        Contact received = IntentCompat', '        ((TextView) findViewById(R.id.detailAvatar))')
snippet('Đổi yêu thích và mở trình quay số', 'ContactDetailActivity.java', '        CheckBox favorite = findViewById', '        findViewById(R.id.deleteButton)')
snippet('Xác nhận xóa khỏi ArrayList', 'ContactDetailActivity.java', '        findViewById(R.id.deleteButton)', None)
p('BaseActivity gắn nút backButton với finish(). Khi xóa, repository.delete(id) loại đối tượng khỏi ArrayList; MainActivity.onResume() cập nhật danh sách và thống kê. Khi thiết bị thiếu trình quay số, ActivityNotFoundException được xử lý bằng Toast.')

page('7 Xóa liên hệ và cập nhật thống kê')
p('Nút Xóa liên hệ mở AlertDialog xác nhận có tên người cần xóa. Chọn Hủy giữ nguyên dữ liệu. Chọn Xóa liên hệ gọi repository.delete(id), hiển thị Toast và finish(). Khi trở lại, MainActivity.onResume() nạp danh sách và đếm lại số yêu thích.')
images('10_delete_confirmation.png','11_contacts_deleted.png','Hình 9 bên trái: hộp thoại xác nhận xóa. Hình 10 bên phải: danh sách sau khi xóa Nguyễn Văn An.')
p('Trong kịch bản kiểm thử, Nguyễn Văn An được đánh dấu yêu thích trước khi xóa: thống kê tăng từ 1 lên 2 yêu thích. Hủy hộp thoại vẫn còn 3 liên hệ; xác nhận xóa còn 2 liên hệ và 1 yêu thích. Đây là dữ liệu mẫu được khởi tạo độc lập cho kịch bản, khác với danh sách 4 liên hệ ở Hình 4.')

page('Mô hình dữ liệu và các thành phần XML')
snippet('Lớp Contact chứa đủ các trường của đề', 'Contact.java', 'public class Contact implements Serializable', '    public Contact(')
snippet('ArrayList và ba liên hệ mẫu', 'ContactRepository.java', '    private final ArrayList<Contact>', '    public static ContactRepository getInstance')
snippet('Đếm số liên hệ yêu thích', 'ContactRepository.java', '    public int getFavoriteCount()', None)
p('Các thuộc tính XML tiêu biểu được trích ở trang kế tiếp. RadioGroup mặc định chọn genderMale; Spinner lấy đủ bốn nhóm từ contact_groups. Tên và điện thoại là bắt buộc, còn ghi chú là tùy chọn.')

page('Trích đoạn XML của các trường nhập liệu')
from lxml import etree
android='{http://schemas.android.com/apk/res/android}'
for filename, view_id, title in [
 ('activity_login.xml','passwordInput','Mật khẩu được ẩn khi nhập'),
 ('activity_add_contact.xml','phoneInput','Điện thoại chỉ cho nhập các chữ số'),
 ('activity_add_contact.xml','genderInput','Giới tính mặc định Nam'),
 ('activity_add_contact.xml','notesInput','Ghi chú nhiều dòng và không bắt buộc')]:
 relative='app/src/main/res/layout/' + filename
 tree=etree.parse(str(base/relative))
 view=next(v for v in tree.getroot().iter() if v.get(android+'id')=='@+id/'+view_id)
 doc.add_heading(title,level=2)
 doc.add_paragraph(relative,style='Caption')
 attrs=['id','inputType','digits','checkedButton','minLines','gravity']
 xml=['<'+view.tag]
 for key in attrs:
  value=view.get(android+key)
  if value is not None: xml.append('    android:'+key+'="'+value+'"')
 xml[-1]+=' />' if view.tag!='RadioGroup' else '>'
 if view.tag=='RadioGroup':
  for child in view:
   xml.append('    <RadioButton android:id="'+child.get(android+'id')+'" />')
  xml.append('</RadioGroup>')
 for line in xml:
  para=doc.add_paragraph();para.paragraph_format.space_after=Pt(0)
  para.paragraph_format.keep_with_next=True
  run=para.add_run(line);run.font.name='DejaVu Sans Mono';run.font.size=Pt(9)
 para.paragraph_format.keep_with_next=False
p('Các đoạn trên trích các thuộc tính liên quan tới yêu cầu; những thuộc tính trình bày như style, kích thước và khoảng cách được lược bỏ để dễ đọc. File XML đầy đủ nằm trong repository.')

page('8 Kiểm thử và hướng dẫn chạy')
p('Môi trường thực thi: máy ảo Android, cấu hình Pixel 7, ảnh hệ thống API 35 (Android 15), màn hình 1080 × 2400. Các ảnh trong báo cáo là ảnh chụp trực tiếp bằng UiAutomation.takeScreenshot() khi chạy Espresso. Dự án dùng Java/XML, minSdk 24, compile/target SDK 37, AGP 9.4.1, Gradle 9.6.0 và Gradle daemon JDK 25.')
table(['Kịch bản đã chạy','Kết quả'],[
 ('Đăng nhập bỏ trống hoặc sai mật khẩu','Hiển thị lỗi; ở lại đăng nhập.'),
 ('Đăng nhập admin / 123456','Mở danh sách; LoginActivity kết thúc.'),
 ('Ba liên hệ mẫu','Tổng 3; yêu thích 1; đúng tên/số/nhóm.'),
 ('Tên trống hoặc điện thoại không hợp lệ','Chặn lưu; setError trên trường nhập.'),
 ('Trùng 0901234567','Chặn lưu; danh sách vẫn 3 liên hệ.'),
 ('Thêm 0966666666 với yêu thích và ghi chú','Tổng 4; yêu thích 2; giữ đủ thông tin.'),
 ('Đổi trạng thái yêu thích rồi quay lại','Danh sách và thống kê cập nhật.'),
 ('Mở quay số','Intent ACTION_DIAL, đúng tel:0901234567.'),
 ('Hủy/xác nhận xóa liên hệ','Hủy giữ dữ liệu; xóa còn 2 và 1 yêu thích.'),
 ('Unit test ContactRepositoryTest','3 test: dữ liệu mẫu, validation, thêm/trùng/yêu thích/xóa.'),
 ('Kiểm thử UI ContactFlowTest','5 test luồng tích hợp; có ảnh chụp minh chứng.')])
h('Cách chạy và kiểm tra lại')
p('Mở repository bằng Android Studio, đồng bộ Gradle, chọn máy ảo/điện thoại API 24 trở lên và Run. Đăng nhập admin / 123456. Chạy các lệnh dưới đây ở thư mục dự án:')
for command in ['./gradlew :app:assembleDebug','./gradlew :app:testDebugUnitTest :app:lintDebug','./gradlew :app:connectedDebugAndroidTest']:
 pp=doc.add_paragraph(command); pp.paragraph_format.space_after=Pt(4)
 for run in pp.runs: run.font.name='DejaVu Sans Mono';run.font.size=Pt(9)
p('Các bài kiểm thử đã chạy trên một ảnh hệ thống máy ảo; chưa kiểm thử trên mọi thiết bị Android. Bộ nhớ danh bạ chỉ tồn tại trong vòng đời tiến trình theo phạm vi đề bài. Mã nguồn, báo cáo và ảnh gốc nằm trong repository; các thư mục sinh tự động và local.properties được loại bằng .gitignore.')

page('Minh chứng kiểm tra dữ liệu nhập')
p('Bên cạnh trạng thái sai tài khoản và số trùng, ứng dụng cũng chặn trường hợp bỏ trống tên đăng nhập hoặc mật khẩu, tên liên hệ trống và số điện thoại sai định dạng. Hai ảnh dưới đây ghi lại thông báo trực tiếp khi thao tác trên máy ảo.')
images('12_login_empty.png','04_add_invalid.png','Hình 11 bên trái: đăng nhập thiếu thông tin. Hình 12 bên phải: số điện thoại chưa đủ 10 chữ số.')
p('Kiểm thử còn nhập số 9012345678 (đủ 10 chữ số nhưng không bắt đầu bằng 0) và 09012345678 (11 chữ số); cả hai bị từ chối và tổng liên hệ không thay đổi. Ghi chú được phép để trống; các nhóm và ba giới tính đã được kiểm tra khi thêm liên hệ.')

page('9 Đối chiếu yêu cầu nộp bài')
p('Bảng dưới ánh xạ các mục của đề tới chức năng, minh chứng và phần báo cáo. Bảng đối chiếu chi tiết nằm trong docs/DOI_CHIEU_YEU_CAU.md của repository.')
table(['Mục của đề','Nội dung và minh chứng'],[
 ('1 Phạm vi','Java/XML; 4 Activity; Intent và Serializable; Contact đủ 6 trường nghiệp vụ; ArrayList trong RAM.'),
 ('2.1 Đăng nhập','EditText tên/mật khẩu ẩn; lỗi trống/sai; admin / 123456; Intent và finish(); hình 1, 2 và mã đăng nhập.'),
 ('2.2 Danh sách','ListView/Adapter đủ avatar, tên, số, nhóm, sao; thống kê; thêm/chi tiết; đúng 3 mẫu; hình 3, 4.'),
 ('2.3 Thêm liên hệ','Đủ 7 thành phần; 4 nhóm; 3 giới tính mặc định Nam; ghi chú tùy chọn; kiểm tra trống/10 số/đầu 0/trùng; hình 5, 6.'),
 ('2.4 Chi tiết','Đủ thông tin; CheckBox yêu thích; ACTION_DIAL; Quay lại; Xóa; cập nhật thống kê; hình 7 đến 10.'),
 ('3.1 Sản phẩm nộp','Dự án Android Studio trên GitHub; báo cáo Word; mã nguồn và kiểm thử; .gitignore loại đúng tệp sinh tự động.'),
 ('3.2 Hình ảnh và mô tả','Cả 4 màn hình có ảnh chụp ứng dụng, chú thích, mô tả component/lớp/hàm/sự kiện và trích đoạn mã.'),
 ('3.3 Định danh','B24DCCC047_PhamHaiDang_BaoCao.docx; đúng tên repository; đầu báo cáo có họ tên, MSSV, D24CQPM03 và link Git.'),
 ('3.4 Minh chứng khớp mã nguồn','Ảnh do Espresso/UiAutomation chụp trên API 35; trích đoạn Java đọc từ file thực tế, kèm tên file và dòng.')])
p('Repository công khai: https://github.com/DANG-PH/B24DCCC047_PhamHaiDang_QuanLyDanhBa. Bài nộp gồm dự án mã nguồn và báo cáo, không đóng gói ZIP. Mỗi kịch bản kiểm thử khởi tạo lại ba liên hệ mẫu; vì vậy ảnh thêm và ảnh xóa minh họa các kịch bản độc lập.')

doc.core_properties.title='Báo cáo ứng dụng quản lý danh bạ liên hệ'
doc.core_properties.author='Phạm Hải Đăng'
doc.core_properties.subject='B24DCCC047 - D24CQPM03'
doc.save(out)
print(out)
